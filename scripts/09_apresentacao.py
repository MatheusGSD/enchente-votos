"""Apresentação em PowerPoint (16:9), separada por tema, com os mapas e os gráficos do projeto.

As figuras saem limpas (só o mapa com a legenda, ou só o gráfico) em saida/apresentacao/; títulos, números,
textos e fontes ficam como texto editável no .pptx. Os mapas usam um estilo só (padrão: pontos).
Uso: python -P 09_apresentacao.py [pontos|areas]
Saída: saida/apresentacao_{estilo}.pptx
"""
import importlib
import json
import sys
from pathlib import Path

import geopandas as gpd
import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

sys.path.insert(0, str(Path(__file__).resolve().parent))
M = importlib.import_module("03_mapas")
MM = importlib.import_module("06_mapas_mudanca")
C = importlib.import_module("07_infograficos")
S = importlib.import_module("08_so_alagados")
COR, SAIDA = M.COR, M.SAIDA
pct, mil, mp, num = M.pct, M.mil, MM.mp, C.num

ESTILO = sys.argv[1] if len(sys.argv) > 1 else "pontos"
FIG = SAIDA / "apresentacao"
DPI_FIG = 200
FONTE = "Segoe UI"
TXT, TXT2, TXT3, FUNDO = "0B0B0B", "52514E", "8A8880", "FCFCFB"
W_SLIDE, H_SLIDE = 13.333, 7.5
# tamanho das figuras = tamanho em que entram no slide (as fontes saem no tamanho real)
MAPA_WH, LARGO_WH, GRAF_WH = (6.8, 5.45), (12.13, 5.0), (7.3, 5.35)
FONTE_2026 = "Fontes: TSE, 1º turno de 2026; UFRGS/IPH, área diretamente afetada em 6 e 8/5/2024."
FONTE_AMBOS = ("Fontes: TSE, 1º turno de 2022 e de 2026 (cada seção comparada com ela mesma); UFRGS/IPH, área "
               "diretamente afetada em 6 e 8/5/2024.")
FONTE_18 = ("Fontes: TSE, 1º turno de 2018, 2022 e 2026 (cada seção comparada com ela mesma); UFRGS/IPH, área "
            "diretamente afetada em 6 e 8/5/2024.")
TEMAS = ["Onde a água chegou", "Voto em 2026", "Mudança no voto", "Eleitorado", "Porto Alegre"]


# ---------- figuras ----------

def salvar(fig, nome):
    FIG.mkdir(exist_ok=True)
    out = FIG / nome
    fig.savefig(out, dpi=DPI_FIG, facecolor=COR["fundo"])
    plt.close(fig)
    return out


def quadro(L, escopo, w_in, h_in):
    if escopo == "mapa":
        return M.ajustar(L["moldura"].total_bounds, w_in, h_in)
    x0, y0, x1, y1 = L["nucleo"][L["nucleo"].NM_MUN == "Porto Alegre"].total_bounds
    return M.ajustar((x0 - 1200, y0 - 1200, x1 + 1200, y1 + 1200), w_in, h_in)


def mancha_poa530(L):
    poa = L["nucleo"][L["nucleo"].NM_MUN == "Porto Alegre"].union_all()
    g = L["mancha_poa530"].union_all().intersection(poa).union(L["mancha_ufrgs"].union_all().difference(poa))
    return gpd.GeoDataFrame(geometry=[g], crs=M.CRS)


def bairros(ax, L, fs=8):
    b = L["bairros"]
    for _, r in b[b.nome.isin(M.BAIRROS_ZOOM)].dissolve("nome").reset_index().iterrows():
        q = r.geometry.representative_point()
        ax.text(q.x, q.y, M.BAIRROS_ZOOM[r.nome], fontsize=fs, color=COR["texto"], ha="center", va="center",
                path_effects=M.HALO, zorder=8)


def camada_grupos(ax, L, smax, vmax):
    """Locais dentro (verde-azulado) e fora (ocre) da área alagada, com as cores dos infográficos."""
    d = L["m_locais"] if ESTILO == "pontos" else L["m_areas"]
    d = d.sort_values("ada")  # os de dentro por cima
    cor = d.ada.map({True: C.COR_ALAGADO, False: C.COR_RESTO})
    if ESTILO == "pontos":
        ax.scatter(d.geometry.x, d.geometry.y, s=M.tamanho(d.validos26, smax, vmax), c=cor,
                   edgecolors=COR["fundo"], linewidths=0.5, zorder=5)
    else:
        d.plot(ax=ax, color=cor, edgecolor=COR["fundo"], lw=0.15, zorder=4)


def amostras(fig, rect, itens, fs=10):
    """Legenda simples: [(tipo, cor, texto)], tipo em ponto, quadrado, mancha, fora, agua."""
    ax = MM.eixo_livre(fig, rect)
    x = 0
    for tipo, cor, txt in itens:
        if tipo == "ponto":
            ax.scatter([x + 0.012], [0.5], s=80, color=cor, clip_on=False)
        elif tipo == "quadrado":
            ax.add_patch(plt.Rectangle((x, 0.3), 0.025, 0.4, color=cor))
        elif tipo == "mancha":
            M.amostra_mancha(ax, x, 0.3, 0.025, 0.4, ESTILO)
        elif tipo == "fora":
            S.amostra_fora(ax, x, 0.3, 0.025, 0.4, ESTILO)
        elif tipo == "agua":
            ax.add_patch(plt.Rectangle((x, 0.3), 0.025, 0.4, facecolor=COR["agua"], lw=0))
        t = ax.text(x + 0.035, 0.5, txt, fontsize=fs, va="center")
        fig.canvas.draw()
        x = ax.transAxes.inverted().transform(t.get_window_extent().corners()[2])[0] + 0.04


def fig_mapa(nome, L, desenhar, legenda, escopo="mapa", paineis=1, titulos=None, wh=None, leg_h=0.15):
    """Um ou dois mapas lado a lado com a legenda embaixo. desenhar(ax, frame, i) pinta o painel i."""
    w, h = wh or (MAPA_WH if paineis == 1 else LARGO_WH)
    fig = plt.figure(figsize=(w, h), dpi=DPI_FIG, facecolor=COR["fundo"])
    topo = 0.94 if titulos else 1
    gap = 0.03
    wp = (1 - gap * (paineis - 1)) / paineis
    for i in range(paineis):
        ax = fig.add_axes([i * (wp + gap), leg_h + 0.01, wp, topo - leg_h - 0.01])
        frame = quadro(L, escopo, w * wp, h * (topo - leg_h - 0.01))
        desenhar(ax, frame, i)
        if titulos:
            ax.set_title(titulos[i], loc="left", fontsize=12, weight="bold", pad=6)
    legenda(fig, [0, 0, 1, leg_h - 0.01])
    return salvar(fig, nome)


def figuras(L, N, NM, N18, NI, vmax):
    f = {}
    smax, smax_poa = 140, 230
    leg_classes = lambda classes, tit, fs=9.5: (lambda fig, r: MM.legenda_classes(
        fig, r, classes, tit, fs, estilo=ESTILO, mancha_x=0.7, lados=M.LADOS if classes is M.CLASSES else None))
    tit26 = "Diferença Lula (PT) − Flávio Bolsonaro (PL) em 2026 (pontos percentuais)"

    def mancha_mun(ax, frame, i):
        M.base(ax, L, frame)
        M.mancha(ax, L["mancha_ufrgs"], ESTILO)
        M.rotulos_municipios(ax, fs=8.5)
        M.rotulo_lago(ax, L, frame, fs=9)
    meio = lambda r: [r[0] + 0.01, r[1] + r[3] * 0.2, r[2] - 0.01, r[3] * 0.6]  # faixa da legenda, centrada
    f["mancha"] = fig_mapa("mapa_mancha.png", L, mancha_mun, lambda fig, r: amostras(
        fig, meio(r), [("mancha", None, "Área alagada em 6 e 8 de maio de 2024"),
                       ("agua", None, "Água que já existia antes")], fs=11), leg_h=0.1)

    def grupos(ax, frame, i):
        M.base(ax, L, frame)
        M.mancha(ax, L["mancha_ufrgs"], ESTILO)
        camada_grupos(ax, L, smax, vmax)
        M.rotulo_lago(ax, L, frame, fs=9)
    tipo = "ponto" if ESTILO == "pontos" else "quadrado"
    f["grupos"] = fig_mapa("mapa_grupos.png", L, grupos, lambda fig, r: amostras(
        fig, meio(r), [(tipo, C.COR_ALAGADO, "Locais dentro da área alagada"),
                       (tipo, C.COR_RESTO, "Locais fora"), ("mancha", None, "Área alagada")], fs=11), leg_h=0.1)

    def voto26(ax, frame, i):
        M.base(ax, L, frame)
        MM.camada(ax, L, ESTILO, "26", smax, vmax)
        M.mancha(ax, L["mancha_ufrgs"], ESTILO)
        M.rotulo_lago(ax, L, frame, fs=9)
    f["voto26"] = fig_mapa("mapa_voto26.png", L, voto26, leg_classes(S.CLASSES_26, tit26))

    def voto26_alag(ax, frame, i):
        S.mapa(ax, L, frame, ESTILO, "26", smax, vmax, fs_lago=9)
    f["voto26_alag"] = fig_mapa("mapa_voto26_alagados.png", L, voto26_alag, lambda fig, r: S.legenda_alagados(
        fig, r, S.CLASSES_26, "Diferença Lula (PT) − Flávio Bolsonaro (PL) em 2026, só nos locais alagados "
        "(pontos percentuais)", 9.5, ESTILO))

    def anos(ax, frame, i):
        M.base(ax, L, frame)
        MM.camada(ax, L, ESTILO, ["22", "26"][i], 95, vmax)
        M.rotulo_lago(ax, L, frame, fs=8)
    f["anos"] = fig_mapa("mapa_2022_2026.png", L, anos, lambda fig, r: MM.legenda_classes(
        fig, [0.2, 0, 0.6, r[3]], MM.CLASSES_MARGEM,
        "Diferença PT − PL (pontos percentuais dos votos válidos)", 9.5, lados=MM.LADOS_PARTIDO), paineis=2,
        titulos=["2022: Lula (PT) × Jair Bolsonaro (PL)", "2026: Lula (PT) × Flávio Bolsonaro (PL)"], leg_h=0.17)

    def mud(ax, frame, i):
        M.base(ax, L, frame)
        MM.camada(ax, L, ESTILO, "mud", smax, vmax)
        M.mancha(ax, L["mancha_ufrgs"], ESTILO)
        M.rotulo_lago(ax, L, frame, fs=9)
    f["mud"] = fig_mapa("mapa_mudanca.png", L, mud, leg_classes(
        MM.CLASSES_MUD, "Mudança na diferença PT − PL de 2022 para 2026 (pontos)"))

    def mud_alag(ax, frame, i):
        S.mapa(ax, L, frame, ESTILO, "mud", smax, vmax, fs_lago=9)
    f["mud_alag"] = fig_mapa("mapa_mudanca_alagados.png", L, mud_alag, lambda fig, r: S.legenda_alagados(
        fig, r, MM.CLASSES_MUD, "Mudança na diferença PT − PL, só nos locais alagados (pontos)", 9.5, ESTILO))

    g530 = mancha_poa530(L)

    def poa_voto(ax, frame, i):
        M.base(ax, L, frame, destaque="Porto Alegre")
        MM.camada(ax, L, ESTILO, "26", smax_poa * 0.75, vmax, municipio="Porto Alegre")
        M.mancha(ax, L["mancha_ufrgs"] if i == 0 else g530, ESTILO)
        bairros(ax, L, fs=7.5)
        M.rotulo_lago(ax, L, frame, fs=8.5)
    f["poa_voto"] = fig_mapa("mapa_poa_voto.png", L, poa_voto, leg_classes(S.CLASSES_26, tit26), escopo="poa",
                             paineis=2,
        titulos=["Mancha de satélite (6 e 8 de maio)", "Simulação do Guaíba a 5,30 m"], leg_h=0.17)

    def poa_mud(ax, frame, i):
        M.base(ax, L, frame, destaque="Porto Alegre")
        MM.camada(ax, L, ESTILO, "mud", smax_poa, vmax, municipio="Porto Alegre")
        M.mancha(ax, L["mancha_ufrgs"], ESTILO)
        bairros(ax, L, fs=8)
        M.rotulo_lago(ax, L, frame, fs=9)
    f["poa_mud"] = fig_mapa("mapa_poa_mudanca.png", L, poa_mud, leg_classes(
        MM.CLASSES_MUD, "Mudança na diferença PT − PL de 2022 para 2026 (pontos)"), escopo="poa")

    # gráficos (as mesmas funções dos infográficos)
    def grafico(nome, desenhar, wh=GRAF_WH):
        fig = plt.figure(figsize=wh, dpi=DPI_FIG, facecolor=COR["fundo"])
        desenhar(fig)
        return salvar(fig, nome)

    f["g_voto"] = grafico("graf_voto.png", lambda fig: C.barras_voto(fig, [0, 0.05, 1, 0.9], N, fs=13),
                          wh=(7.3, 4.2))
    f["g_inclinacao"] = grafico("graf_inclinacao.png",
                                lambda fig: C.inclinacao(fig, [0.06, 0.13, 0.86, 0.84], NM, fs=13))

    def criterios(fig):
        C.legenda_grupos(fig, 0.01, 0.95, 11)
        linhas = [(r, det, s["dentro"]["mudanca_margem"], s["fora"]["mudanca_margem"])
                  for r, det, s in C.linhas_criterios(NM)]
        C.halteres(fig, [0.01, 0.06, 0.98, 0.8], linhas, (-12, 0), [-12, -8, -4, 0], fs=11, rot_quebra=40)
    f["g_criterios"] = grafico("graf_criterios.png", criterios)

    def cidades(fig, chave, xlim, ticks, unidade="", det=None):
        C.legenda_grupos(fig, 0.01, 0.955, 11, resto="restante da mesma cidade")
        linhas = [(k, det(d) if det else None, d[chave], fo[chave]) for k, d, fo in C.cidades(NM)]
        C.halteres(fig, [0.01, 0.05, 0.98, 0.84], linhas, xlim, ticks, fs=11, unidade=unidade)
    f["g_cidades"] = grafico("graf_cidades.png", lambda fig: cidades(
        fig, "mudanca_margem", (-18, 0), [-15, -10, -5, 0],
        det=lambda d: f"{C.mil_k(d['aptos_26'])} eleitores em locais alagados"), wh=(7.3, 5.6))
    f["g_aptos_cidades"] = grafico("graf_aptos_cidades.png", lambda fig: cidades(
        fig, "mudanca_aptos_pct", (-10, 10), [-10, -5, 0, 5, 10], unidade="%",
        det=lambda d: f"{C.saldo(d['aptos_26'] - d['aptos_22'])} eleitores nos locais alagados"), wh=(7.3, 5.6))

    def eleitorado(fig):
        fig.text(0.01, 0.97, "Eleitores aptos, 2022 → 2026", fontsize=12.5, weight="bold", va="top")
        C.barras_aptos(fig, [0.01, 0.53, 0.98, 0.36], NM, fs=12)
        fig.text(0.01, 0.43, "Comparecimento", fontsize=12.5, weight="bold", va="top")
        C.comparecimento(fig, [0.01, 0.05, 0.98, 0.3], NM, fs=11.5)
    f["g_eleitorado"] = grafico("graf_eleitorado.png", eleitorado)
    f["g_decomp"] = grafico("graf_decomp.png", lambda fig: decomposicao(fig, NM), wh=(7.3, 4.6))
    f["g_tres"] = grafico("graf_tres_eleicoes.png", lambda fig: C.tres_eleicoes(fig, [0.01, 0.37, 0.97, 0.57], N18,
                                                                                    fs=13, xmin=-0.8))

    def antes_depois(fig):
        C.legenda_grupos(fig, 0.01, 0.955, 11)
        fig.text(0.01, 0.87, "Eleitores aptos, variação em %", fontsize=12.5, weight="bold", va="top")
        C.antes_depois(fig, [0.01, 0.5, 0.98, 0.33], N18, "mudanca_aptos_pct", fs=12, unidade="%", rot_w=0.3)
        fig.text(0.01, 0.43, "Abstenção, mudança em pontos percentuais", fontsize=12.5, weight="bold", va="top")
        C.antes_depois(fig, [0.01, 0.04, 0.98, 0.33], N18, "mudanca_abstencao", fs=12, rot_w=0.3)
    f["g_antes_depois"] = grafico("graf_antes_depois.png", antes_depois)
    f["g_incerteza"] = grafico("graf_incerteza.png", lambda fig: intervalos(fig, NI))
    return f


MEDIDAS_INC = [("voto", "Diferença PT − PSL/PL", ""), ("aptos", "Eleitores aptos, variação", "%"),
               ("abstencao", "Abstenção", "")]
PERIODOS_INC = [("18_22", "2018 → 2022", "antes da enchente"), ("22_26", "2022 → 2026", "depois da enchente")]


def mp0(x):
    """mp() sem o "−0,0"."""
    return "0,0" if abs(x) < 0.05 else mp(x)


def fora_da_margem(m):
    lo, hi = m["ic95"]
    return lo > 0 or hi < 0


def intervalos(fig, NI, fs=12):
    """Uma faixa por medida: estimativa (alagado − restante) e intervalo de 95%, antes e depois da enchente.
    Cada faixa tem a própria escala (unidades diferentes); o zero é a linha de referência."""
    for j, (k, titulo, u) in enumerate(MEDIDAS_INC):
        topo = 0.99 - j * 0.335
        fig.text(0.01, topo, titulo + (", em %" if u else ", em pontos"), fontsize=fs + 0.5, weight="bold",
                 va="top")
        ms = [NI["medidas"][f"{k}_{p}"] for p, _, _ in PERIODOS_INC]
        lim = max(abs(v) for m in ms for v in m["ic95"]) * 1.15
        ax = fig.add_axes([0.3, topo - 0.29, 0.36, 0.22])
        ax.set_xlim(-lim, lim)
        ax.set_ylim(1.6, -0.6)
        ax.axis("off")
        ax.axvline(0, color=COR["texto2"], lw=1, zorder=1)
        tr = ax.get_yaxis_transform()
        for i, ((_, per, quando), m) in enumerate(zip(PERIODOS_INC, ms)):
            lo, hi = m["ic95"]
            ax.plot([lo, hi], [i, i], color=COR["texto3"], lw=3, solid_capstyle="round", zorder=2)
            ax.scatter([m["estimativa"]], [i], s=70, color=COR["texto"], zorder=3)
            ax.text(-0.3 / 0.36 + 0.01 / 0.36, i - 0.08, per, transform=tr, va="bottom", fontsize=fs, weight="bold")
            ax.text(-0.3 / 0.36 + 0.01 / 0.36, i - 0.02, quando, transform=tr, va="top", fontsize=fs - 1.5,
                    color=COR["texto2"])
            neg = "bold" if fora_da_margem(m) else "normal"
            ax.text(1.05, i - 0.08, f"{mp0(m['estimativa'])}{u}", transform=tr, va="bottom", fontsize=fs + 1,
                    weight=neg)
            ax.text(1.05, i - 0.02, f"de {mp0(lo)}{u} a {mp0(hi)}{u}", transform=tr, va="top", fontsize=fs - 1.5,
                    color=COR["texto2"])


def decomposicao(fig, NM):
    """Barras: quanto cada parte tirou dos votos válidos, nos alagados e no restante (log × 100)."""
    d, f = NM["criterios"]["ada"]["dentro"], NM["criterios"]["ada"]["fora"]
    partes = [("decomp_aptos", "Menos eleitores aptos"), ("decomp_comparecimento", "Mais abstenção"),
              ("decomp_brancos_nulos", "Mais brancos e nulos")]
    C.legenda_grupos(fig, 0.01, 0.94, 11)
    ax = fig.add_axes([0.33, 0.08, 0.5, 0.76])
    vmax = max(-d[k] for k, _ in partes)
    ax.set_xlim(0, vmax * 1.12)
    ax.set_ylim(len(partes) - 0.4, -0.6)
    ax.axis("off")
    for i, (k, rot) in enumerate(partes):
        for dy, s, cor in [(-0.17, d, C.COR_ALAGADO), (0.17, f, C.COR_RESTO)]:
            ax.barh(i + dy, -s[k], height=0.3, color=cor)
            ax.text(-s[k] + vmax * 0.02, i + dy, mp(s[k]), va="center", fontsize=12, weight="bold")
        ax.text(-0.66, i, rot, transform=ax.get_yaxis_transform(), va="center", fontsize=12.5)
    ax.axvline(0, color=COR["texto3"], lw=1)


# ---------- slides ----------

class Deck:
    def __init__(self):
        self.prs = Presentation()
        self.prs.slide_width, self.prs.slide_height = Inches(W_SLIDE), Inches(H_SLIDE)
        self.n = 0

    def novo(self, tema=None, titulo=None, fonte=None):
        s = self.prs.slides.add_slide(self.prs.slide_layouts[6])
        s.background.fill.solid()
        s.background.fill.fore_color.rgb = RGBColor.from_string(FUNDO)
        self.n += 1
        if tema:
            texto(s, 0.6, 0.38, 9, 0.3, [par(tema.upper(), 11, TXT2, negrito=True)])
        if titulo:
            texto(s, 0.6, 0.66, 12.1, 0.8, [par(titulo, 28, TXT, negrito=True)])
        if fonte:
            texto(s, 0.6, 7.0, 10, 0.35, [par(fonte, 9, TXT3)])
        if self.n > 1:
            texto(s, 10.73, 7.0, 2.0, 0.3, [par(f"Matheus Dalmolin · {self.n}", 9, TXT2)], alinhar=PP_ALIGN.RIGHT)
        return s


def par(t, tam=16, cor=TXT, negrito=False, antes=0, espaco=1.15):
    """t: texto ou lista de (trecho, negrito) para misturar peso na mesma linha."""
    return dict(t=t, tam=tam, cor=cor, negrito=negrito, antes=antes, espaco=espaco)


def texto(s, x, y, w, h, pars, alinhar=None, ancora=MSO_ANCHOR.TOP):
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = ancora
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    for i, p in enumerate(pars):
        pp = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        pp.space_before = Pt(p["antes"])
        pp.line_spacing = p["espaco"]
        if alinhar:
            pp.alignment = alinhar
        trechos = p["t"] if isinstance(p["t"], list) else [(p["t"], p["negrito"])]
        for t, neg in trechos:
            r = pp.add_run()
            r.text = t
            r.font.name, r.font.size, r.font.bold = FONTE, Pt(p["tam"]), neg
            r.font.color.rgb = RGBColor.from_string(p["cor"])
    return tb


def bolinha(s, x, y, d, cor):
    o = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(d), Inches(d))
    o.fill.solid()
    o.fill.fore_color.rgb = RGBColor.from_string(cor.lstrip("#").upper())
    o.line.fill.background()


def numero(s, x, y, valor, rotulo, cor=None, tam=36, w=2.4):
    """Número grande com rótulo embaixo e, se houver cor, a bolinha do grupo à esquerda."""
    dx = 0
    if cor:
        bolinha(s, x, y + tam / 72 * 0.42, 0.2, cor)
        dx = 0.32
    texto(s, x + dx, y, w - dx, tam / 72 * 1.25, [par(valor, tam, TXT, negrito=True, espaco=1.0)])
    texto(s, x + dx, y + tam / 72 * 1.3, w - dx, 0.6, [par(rotulo, 12.5, TXT2, espaco=1.1)])


def imagem(s, arq, x, y, w=None, h=None):
    s.shapes.add_picture(str(arq), Inches(x), Inches(y), width=Inches(w) if w else None,
                         height=Inches(h) if h else None)


def slide_mapa(d, tema, titulo, arq, pars, numeros=(), fonte=FONTE_2026):
    s = d.novo(tema, titulo, fonte)
    w, h = MAPA_WH
    imagem(s, arq, W_SLIDE - 0.5 - w, 1.45, w=w)
    y = 1.6
    for valor, rot, cor in numeros:
        numero(s, 0.6, y, valor, rot, cor, w=5.2)
        y += 1.25
    texto(s, 0.6, y + 0.1, 5.3, 6.8 - y, pars)
    return s


def slide_grafico(d, tema, titulo, arq, pars, fonte=FONTE_AMBOS, wh=GRAF_WH):
    s = d.novo(tema, titulo, fonte)
    w, h = wh
    imagem(s, arq, W_SLIDE - 0.5 - w, 1.5 + (5.35 - h) / 2, w=w)
    texto(s, 0.6, 1.65, W_SLIDE - 1.6 - w, 5.2, pars)
    return s


def slide_largo(d, tema, titulo, arq, lead, fonte):
    """Figura larga embaixo de um texto de até duas linhas; a figura encolhe para caber até o rodapé."""
    s = d.novo(tema, titulo, fonte)
    texto(s, 0.6, 1.35, 12.1, 0.6, [par(lead, 14, TXT2, espaco=1.15)])
    y, h = 2.05, 6.9 - 2.05
    w = h * LARGO_WH[0] / LARGO_WH[1]
    imagem(s, arq, (W_SLIDE - w) / 2, y, h=h)
    return s


def divisor(d, i, subtitulo):
    s = d.novo()
    r = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.6), Inches(3.05), Inches(1.0), Inches(0.07))
    r.fill.solid()
    r.fill.fore_color.rgb = RGBColor.from_string(TXT)
    r.line.fill.background()
    texto(s, 0.6, 1.55, 3, 1.4, [par(str(i), 80, TXT, negrito=True, espaco=1.0)])
    texto(s, 0.6, 3.35, 11, 1.0, [par(TEMAS[i - 1], 40, TXT, negrito=True)])
    texto(s, 0.6, 4.35, 10, 1.5, [par(subtitulo, 18, TXT2, espaco=1.3)])
    return s


def bullets(itens, tam=16, antes=10):
    return [par(t, tam, TXT, antes=0 if i == 0 else antes, espaco=1.2) for i, t in enumerate(itens)]


def tabela_secoes(s, NM, x, y, w):
    lst = NM["secoes_alagadas_que_mudaram_de_local"]
    cab = ["Cidade", "Local em 2022", "Local em 2026", "Seções", "Eleitores", "Distância"]
    larg = [1.7, 3.3, 3.3, 0.9, 1.1, 1.13]
    linhas = [[x_["municipio"], C.nome_curto(x_["nome_2022"]), C.nome_curto(x_["nome_2026"]), str(x_["secoes"]),
               mil(x_["aptos_2022"]),
               "—" if x_["distancia_m"] is None else (f"{num(x_['distancia_m'] / 1000)} km"
                                                       if x_["distancia_m"] >= 1000 else f"{x_['distancia_m']} m")]
              for x_ in lst]
    linhas.append(["Total", "", "", str(sum(x_["secoes"] for x_ in lst)), mil(sum(x_["aptos_2022"] for x_ in lst)), ""])
    gf = s.shapes.add_table(len(linhas) + 1, len(cab), Inches(x), Inches(y), Inches(w), Inches(0.36 * (len(linhas) + 1)))
    t = gf.table
    for j, lw in enumerate(larg):
        t.columns[j].width = Inches(lw)
    for i, row in enumerate([cab] + linhas):
        t.rows[i].height = Inches(0.36)
        for j, v in enumerate(row):
            c = t.cell(i, j)
            c.fill.solid()
            c.fill.fore_color.rgb = RGBColor.from_string("EFEDE8" if i == 0 else (FUNDO if i % 2 else "F6F5F2"))
            c.margin_top = c.margin_bottom = Inches(0.03)
            c.vertical_anchor = MSO_ANCHOR.MIDDLE
            tf = c.text_frame
            tf.text = ""
            p = tf.paragraphs[0]
            p.alignment = PP_ALIGN.RIGHT if j >= 3 else PP_ALIGN.LEFT
            r = p.add_run()
            r.text = v
            r.font.name, r.font.size = FONTE, Pt(12)
            r.font.bold = i == 0 or i == len(linhas) or j == 2
            r.font.color.rgb = RGBColor.from_string(TXT2 if i == 0 else TXT)


def peso_entorno(L):
    """% dos moradores do entorno que estavam na mancha, nos locais de dentro e de fora (média pelos aptos)."""
    pts = L["locais"]
    f = pd.read_csv(SAIDA / "peso_alagado.csv", index_col="pid").f
    p = pts.groupby("pid").agg(ada=("ada", "first"), aptos=("aptos", "sum"))
    p["f"] = f.reindex(p.index).fillna(0)
    return {k: 100 * np.average(g.f, weights=g.aptos) for k, g in p.groupby("ada")}


def montar(L, N, NM, N18, NI, F):
    d = Deck()
    ada, ada_m = N["criterios"]["ada"], NM["criterios"]["ada"]
    dn, fo = ada["dentro"], ada["fora"]
    md, mf = ada_m["dentro"], ada_m["fora"]
    pz = N["ponderado_populacao"]
    entorno = peso_entorno(L)
    # capa
    s = d.novo()
    texto(s, 0.6, 1.6, 5.2, 2.8, [par("Enchente de 2024 e voto para presidente na Grande Porto Alegre", 34, TXT,
                                     negrito=True, espaco=1.05)])
    texto(s, 0.6, 4.45, 5.2, 1.2, [par("Locais de votação dentro e fora da área alagada, no 1º turno de 2022 e de "
                                       "2026", 17, TXT2, espaco=1.25)])
    texto(s, 0.6, 5.7, 5.2, 0.4, [par("Matheus Dalmolin · outubro de 2026", 14, TXT2)])
    imagem(s, F["grupos"], W_SLIDE - 0.5 - MAPA_WH[0], 1.0, w=MAPA_WH[0])

    # o que foi comparado
    s = d.novo("Antes de começar", "O que foi comparado", "Comparação descritiva: mostra onde o voto e o eleitorado "
                                                            "mudaram diferente, não mede o efeito da enchente.")
    texto(s, 0.6, 1.65, 7.2, 5.2, bullets([
        [("Unidade: ", True), (f"o local de votação. São {mil(N['total']['locais'])} locais em "
                               f"{N['municipios_no_mapa']} municípios, inteiros ou em parte, num retângulo em volta "
                               "da Grande Porto Alegre.", False)],
        [("Área alagada: ", True), ("mancha da UFRGS/IPH com as imagens de 6 e 8 de maio de 2024, sem a água que já "
                                    "existia. Um local conta como alagado quando o ponto dele fica dentro da mancha.",
                                    False)],
        [("Voto: ", True), ("votos válidos para presidente no 1º turno. Lula (PT) × Flávio Bolsonaro (PL) em 2026; "
                            "Lula (PT) × Jair Bolsonaro (PL) em 2022.", False)],
        [("2022 → 2026: ", True), ("cada seção eleitoral é comparada com ela mesma nas duas eleições, no local onde "
                                   "votou em 2026.", False)],
        [("Antes da enchente: ", True), ("as mesmas seções em 2018, Fernando Haddad (PT) × Jair Bolsonaro (PSL).",
                                         False)],
    ], tam=16, antes=14))
    texto(s, 8.6, 1.65, 4.2, 4.5, [par("Temas", 13, TXT2, negrito=True)] +
          [par(f"{i}  {t}", 18, TXT, antes=10) for i, t in enumerate(TEMAS, 1)])

    # 1. onde a água chegou
    divisor(d, 1, "A área alagada em maio de 2024 e quais locais de votação ficaram dentro dela")
    slide_mapa(d, TEMAS[0], "A área diretamente afetada em maio de 2024", F["mancha"], [
        par("Mancha da UFRGS/IPH com as imagens de satélite de 6 e 8 de maio de 2024. A água que já existia (Guaíba, "
            "rios e delta do Jacuí) não conta como alagada.", 15, TXT2, espaco=1.25),
        par("A imagem de satélite deixa passar a água entre prédios, como no Centro de Porto Alegre; por isso os "
            "números foram refeitos com outras manchas e com folga de 500 m.", 15, TXT2, antes=12, espaco=1.25)],
        numeros=[(f"{mil(pz['pop_alagada'] / 1000)} mil", f"moradores dentro da mancha (Censo 2022), de "
                                                          f"{num(pz['pop_total'] / 1e6)} milhões na área do mapa",
                  None)],
        fonte="Fontes: UFRGS/IPH, Banco de dados das cheias da RH do Guaíba v1.4 (CC BY 4.0); IBGE, Censo 2022; "
              "SEMA-RS, base cartográfica 1:25.000.")
    slide_mapa(d, TEMAS[0], "Quais locais de votação contam como alagados", F["grupos"], [
        par(f"Nos locais alagados, {num(entorno[True])}% dos moradores do entorno estavam dentro da mancha; nos "
            f"outros, {num(entorno[False])}%.", 15, TXT2, espaco=1.25)],
        numeros=[(f"{dn['locais']} locais", "de votação com o ponto dentro da área alagada", C.COR_ALAGADO),
                 (f"{mil(dn['aptos'] / 1000)} mil", f"eleitores votam neles, "
                                                   f"{pct(100 * dn['aptos'] / N['aptos_mapa_total'])} do eleitorado da "
                                                   "área do mapa", None)],
        fonte="Fontes: TSE, locais de votação de 2026; UFRGS/IPH, área diretamente afetada em 6 e 8/5/2024; "
              "IBGE, Censo 2022.")

    # 2. voto em 2026
    divisor(d, 2, "Como votaram os locais alagados no 1º turno de 2026, comparados com o restante")
    slide_mapa(d, TEMAS[1], "Voto em 2026 e área alagada", F["voto26"], [
        par("Cada círculo é um local de votação; o tamanho é o número de votos válidos." if ESTILO == "pontos" else
            "Cada área é o entorno de um local de votação, onde moram 100 ou mais pessoas por km².", 15, TXT2,
            espaco=1.25)],
        numeros=[(f"{pct(dn['pct_lula'])} × {pct(dn['pct_flavio'])}", "Lula × Flávio Bolsonaro nos locais alagados",
                  C.COR_ALAGADO),
                 (f"{pct(fo['pct_lula'])} × {pct(fo['pct_flavio'])}", "Lula × Flávio Bolsonaro no restante", C.COR_RESTO)])
    slide_mapa(d, TEMAS[1], "Só os locais alagados", F["voto26_alag"], [
        par("Mesmo mapa, pintando só os locais cujo ponto fica dentro da área alagada. Os outros aparecem em contorno "
            "cinza.", 15, TXT2, espaco=1.25),
        par("Os locais alagados se concentram em Canoas, Porto Alegre, São Leopoldo, Eldorado do Sul e Guaíba.",
            15, TXT2, antes=12, espaco=1.25)],
        numeros=[(mp(dn["pct_lula"] - dn["pct_flavio"]), "diferença Lula − Flávio Bolsonaro nos locais alagados (pontos)",
                  C.COR_ALAGADO),
                 (mp(fo["pct_lula"] - fo["pct_flavio"]), "no restante", C.COR_RESTO)])
    slide_grafico(d, TEMAS[1], "Locais alagados × restante", F["g_voto"], [
        par("% dos votos válidos. Em cinza, os outros candidatos somados.", 15, TXT2, espaco=1.25),
        par(f"Os locais alagados votaram um pouco mais em Lula. Mas nos mesmos lugares, em 2022, a diferença já "
            f"existia: Lula − Jair Bolsonaro ficou em {mp(md['margem_22'])} nos alagados e {mp(mf['margem_22'])} no "
            "restante.", 15, TXT, antes=14, espaco=1.25)], fonte=FONTE_AMBOS, wh=(7.3, 4.2))

    # 3. mudança
    divisor(d, 3, "O que mudou de 2022 para 2026 nos mesmos lugares, dentro e fora da área alagada, e como era "
                  "em 2018")
    slide_largo(d, TEMAS[2], "2022 e 2026 nos mesmos lugares", F["anos"],
                "Em 2022, cada seção eleitoral aparece no local onde votou em 2026. O candidato do PL foi Jair "
                "Bolsonaro em 2022 e Flávio Bolsonaro em 2026.", FONTE_AMBOS)
    slide_mapa(d, TEMAS[2], "Mudança na diferença PT − PL", F["mud"], [
        par("Quase todos os locais andaram para o PL, dentro e fora da área alagada.", 15, TXT2, espaco=1.25)],
        numeros=[(f"{mp(md['mudanca_margem'])} pontos", "nos locais alagados", C.COR_ALAGADO),
                 (f"{mp(mf['mudanca_margem'])} pontos", "no restante", C.COR_RESTO)], fonte=FONTE_AMBOS)
    slide_mapa(d, TEMAS[2], "Mudança: só os locais alagados", F["mud_alag"], [
        par("Mesma mudança, pintando só os locais dentro da área alagada.", 15, TXT2, espaco=1.25)],
        numeros=[(f"{mp(md['mudanca_margem'] - mf['mudanca_margem'])} ponto",
                  "diferença entre a mudança nos alagados e no restante", None)], fonte=FONTE_AMBOS)
    slide_grafico(d, TEMAS[2], "As duas linhas andaram juntas", F["g_inclinacao"], [
        par("Diferença PT − PL, em pontos percentuais dos votos válidos.", 15, TXT2, espaco=1.25),
        par([("Lula (PT): ", True), (f"{pct(md['pct_lula_22'])} → {pct(md['pct_lula_26'])} nos alagados, "
                                f"{pct(mf['pct_lula_22'])} → {pct(mf['pct_lula_26'])} no restante.", False)],
            15, TXT, antes=16, espaco=1.25),
        par([("PL (Jair → Flávio Bolsonaro): ", True), (f"{pct(md['pct_n22_22'])} → {pct(md['pct_n22_26'])} nos alagados, "
                              f"{pct(mf['pct_n22_22'])} → {pct(mf['pct_n22_26'])} no restante.", False)],
            15, TXT, antes=10, espaco=1.25)])
    t18d, t18f = N18["criterios"]["ada"]["dentro"], N18["criterios"]["ada"]["fora"]
    dist18 = [t18d[f"margem_{a}"] - t18f[f"margem_{a}"] for a, _, _ in C.ANOS_18]
    slide_grafico(d, TEMAS[2], "E antes da enchente? 2018, 2022 e 2026", F["g_tres"], [
        par("Diferença PT − PSL/PL nas seções que existem nas três eleições. Em 2018, a disputa foi Fernando Haddad "
            "(PT) × Jair Bolsonaro (PSL).", 15, TXT2, espaco=1.25),
        par(f"A distância entre os locais alagados e o restante ficou entre {mp(min(dist18))} e {mp(max(dist18))} "
            "pontos nas três eleições, antes e depois da enchente.", 15, TXT, antes=14, espaco=1.25),
        par(C.NOTA_CENARIO, 13, TXT2, antes=14, espaco=1.25)], fonte=FONTE_18)
    difs = [s_["dentro"]["mudanca_margem"] - s_["fora"]["mudanca_margem"] for _, _, s_ in C.linhas_criterios(NM)]
    slide_grafico(d, TEMAS[2], "Com outras definições de área alagada", F["g_criterios"], [
        par("Mudança na diferença PT − PL, em pontos, com seis jeitos de definir \"local alagado\".", 15, TXT2,
            espaco=1.25),
        par(f"A diferença entre alagados e restante fica entre {mp(min(difs))} e {mp(max(difs))} ponto: o resultado "
            "não depende da mancha escolhida.", 15, TXT, antes=14, espaco=1.25)],
        fonte=FONTE_AMBOS[:-1] + "; Governo do RS/SPGG, área inundada em 23/5/2024; IBGE, Censo 2022.")
    cs = C.cidades(NM)
    mais = sum(a["mudanca_margem"] < b["mudanca_margem"] for _, a, b in cs)
    slide_grafico(d, TEMAS[2], "Dentro de cada cidade", F["g_cidades"], [
        par(f"Mudança na diferença PT − PL, em pontos. Cidades com {mil(C.MIN_APTOS_CIDADE)} ou mais eleitores em "
            "locais alagados.", 15, TXT2, espaco=1.25),
        par(f"Em {mais} das {len(cs)} cidades, os locais alagados andaram mais para o PL que o resto da cidade; nas "
            f"outras {len(cs) - mais}, menos. Não há um padrão.", 15, TXT, antes=14, espaco=1.25)], wh=(7.3, 5.6))

    # 4. eleitorado
    divisor(d, 4, "Eleitores aptos, abstenção, brancos e nulos: o que mudou mais nos locais alagados")
    slide_grafico(d, TEMAS[3], "Eleitores aptos e comparecimento", F["g_eleitorado"], [
        par("Só seções que existem nas duas eleições.", 15, TXT2, espaco=1.25),
        par(f"Os locais alagados perderam {num(-md['mudanca_aptos_pct'])}% dos eleitores aptos; o restante, "
            f"{num(-mf['mudanca_aptos_pct'])}%. A abstenção subiu {num(md['mudanca_abstencao'])} pontos nos alagados "
            f"e {num(mf['mudanca_abstencao'])} no restante.", 15, TXT, antes=14, espaco=1.25)])
    slide_grafico(d, TEMAS[3], "Eleitores antes e depois da enchente", F["g_antes_depois"], [
        par("Seções que existem em 2018, 2022 e 2026.", 15, TXT2, espaco=1.25),
        par(f"De 2018 para 2022, os dois grupos ganharam eleitores no mesmo ritmo ({mp(t18d['mudanca_aptos_pct_18_22'])}% "
            f"e {mp(t18f['mudanca_aptos_pct_18_22'])}%), e a abstenção subiu praticamente o mesmo.", 15, TXT, antes=14,
            espaco=1.25),
        par(f"De 2022 para 2026, os locais alagados perderam {num(-t18d['mudanca_aptos_pct_22_26'])}% dos eleitores, "
            f"contra {num(-t18f['mudanca_aptos_pct_22_26'])}% no restante.", 15, TXT, antes=10, espaco=1.25),
        par(C.texto_cidades_18(N18), 13, TXT2, antes=14, espaco=1.25)], fonte=FONTE_18)
    gap = (md["decomp_aptos"] + md["decomp_comparecimento"] + md["decomp_brancos_nulos"]) - \
          (mf["decomp_aptos"] + mf["decomp_comparecimento"] + mf["decomp_brancos_nulos"])
    parte = lambda k: 100 * (md[k] - mf[k]) / gap
    slide_grafico(d, TEMAS[3], "De onde vêm os votos válidos a menos", F["g_decomp"], [
        par(f"Votos válidos das mesmas seções: {mp(md['mudanca_validos_pct'])}% nos alagados e "
            f"{mp(mf['mudanca_validos_pct'])}% no restante.", 15, TXT2, espaco=1.25),
        par(f"Da perda a mais nos alagados, cerca de {parte('decomp_aptos'):.0f}% vem da queda de eleitores aptos e "
            f"{parte('decomp_comparecimento'):.0f}% da abstenção. Brancos e nulos subiram igual nos dois grupos "
            f"(brancos {mp(md['mudanca_brancos'])} e {mp(mf['mudanca_brancos'])} ponto).", 15, TXT, antes=14,
            espaco=1.25),
        par("Os dados não separam as causas: mudança de endereço, mortes, novos eleitores, remanejamento de seções.",
            13, TXT2, antes=14, espaco=1.25)], wh=(7.3, 4.6))
    menos = sum(a["mudanca_aptos_pct"] < b["mudanca_aptos_pct"] for _, a, b in cs)
    slide_grafico(d, TEMAS[3], "Eleitores aptos em cada cidade", F["g_aptos_cidades"], [
        par("Variação de 2022 para 2026, em %, nos locais alagados e no resto da mesma cidade.", 15, TXT2,
            espaco=1.25),
        par(f"Em {menos} das {len(cs)} cidades, os locais alagados perderam mais eleitores (ou ganharam menos) que o "
            "resto da cidade.", 15, TXT, antes=14, espaco=1.25)], wh=(7.3, 5.6))
    s = d.novo(TEMAS[3], "Seções de locais alagados que mudaram de local", FONTE_AMBOS)
    texto(s, 0.6, 1.4, 12.1, 0.5, [par("Seções que em 2022 votavam num local dentro da área alagada e em 2026 votaram "
                                       "em outro, a mais de 250 m. — = local de 2026 sem coordenada no TSE.", 14, TXT2)])
    tabela_secoes(s, NM, 0.6, 2.0, 11.43)

    # 5. Porto Alegre
    pa, p5 = N["porto_alegre_criterios"]["ada"], N["porto_alegre_criterios"]["poa530"]
    pm = NM["porto_alegre_criterios"]["ada"]
    divisor(d, 5, "A capital, onde a mancha de satélite subestima a água no Centro")
    slide_largo(d, TEMAS[4], "Porto Alegre: voto em 2026 com duas manchas", F["poa_voto"],
                f"Locais alagados pela mancha de satélite ({pa['dentro']['locais']}): Lula "
                f"{pct(pa['dentro']['pct_lula'])} × Flávio Bolsonaro {pct(pa['dentro']['pct_flavio'])}. Pela simulação de "
                f"5,30 m ({p5['dentro']['locais']}): {pct(p5['dentro']['pct_lula'])} × "
                f"{pct(p5['dentro']['pct_flavio'])}. Restante: {pct(pa['fora']['pct_lula'])} × "
                f"{pct(pa['fora']['pct_flavio'])}.",
                "Fontes: TSE, 1º turno de 2026; UFRGS/IPH, área diretamente afetada em 6 e 8/5/2024 e simulação "
                "hidrostática do Guaíba a 5,30 m.")
    slide_mapa(d, TEMAS[4], "Porto Alegre: mudança de 2022 para 2026", F["poa_mud"], [
        par("No Centro Histórico, na Cidade Baixa e no Menino Deus, vários locais andaram para o PT.", 15, TXT2,
            espaco=1.25),
        par(f"Eleitores aptos: {mp(pm['dentro']['mudanca_aptos_pct'])}% nos alagados e "
            f"{mp(pm['fora']['mudanca_aptos_pct'])}% no restante. Abstenção: {mp(pm['dentro']['mudanca_abstencao'])} "
            f"e {mp(pm['fora']['mudanca_abstencao'])} ponto.", 15, TXT, antes=12, espaco=1.25)],
        numeros=[(f"{mp(pm['dentro']['mudanca_margem'])} pontos", "nos locais alagados", C.COR_ALAGADO),
                 (f"{mp(pm['fora']['mudanca_margem'])} pontos", "no restante da cidade", C.COR_RESTO)],
        fonte=FONTE_AMBOS)

    # fechamento
    mi = NI["medidas"]
    faixa = lambda m, u="": f"de {mp0(m['ic95'][0])}{u} a {mp0(m['ic95'][1])}{u}"
    margem = lambda m: ("maior que a margem de acaso" if fora_da_margem(m) else "dentro da margem de acaso")
    antes_fora = [rot for k, rot, _ in MEDIDAS_INC if fora_da_margem(mi[f"{k}_18_22"])]
    s_ = d.novo("Fechamento", "Quanto pode ser acaso", FONTE_18[:-1] + "; intervalo de 95% por reamostragem das "
                f"{NI['zonas']} zonas eleitorais ({mil(NI['b'])} sorteios).")
    imagem(s_, F["g_incerteza"], W_SLIDE - 0.5 - GRAF_WH[0], 1.5, w=GRAF_WH[0])
    texto(s_, 0.6, 1.65, W_SLIDE - 1.6 - GRAF_WH[0], 5.2, [
        par("Diferença entre a mudança nos locais alagados e no restante, com o intervalo de 95%: a faixa onde a "
            "diferença cai quando as zonas eleitorais são sorteadas de novo. Se a faixa cruza o zero, pode ser acaso.",
            13.5, TXT2, espaco=1.25),
        par([("Voto: ", True), (f"{mp0(mi['voto_22_26']['estimativa'])} ponto, {faixa(mi['voto_22_26'])}; "
                                f"{margem(mi['voto_22_26'])}.", False)], 15, TXT, antes=14, espaco=1.25),
        par([("Eleitores aptos: ", True), (f"{mp0(mi['aptos_22_26']['estimativa'])}%, {faixa(mi['aptos_22_26'], '%')}; "
                                           f"{margem(mi['aptos_22_26'])}.", False)], 15, TXT, antes=10, espaco=1.25),
        par([("Abstenção: ", True), (f"{mp0(mi['abstencao_22_26']['estimativa'])} ponto, "
                                     f"{faixa(mi['abstencao_22_26'])}; {margem(mi['abstencao_22_26'])}.", False)],
            15, TXT, antes=10, espaco=1.25),
        par("Antes da enchente (2018 → 2022), " + ("as três diferenças ficam dentro da margem." if not antes_fora else
                                                   "ficam fora da margem: " + ", ".join(antes_fora) + "."),
            13.5, TXT2, antes=14, espaco=1.25)])
    s = d.novo("Fechamento", "O que os dados mostram", FONTE_18)
    texto(s, 0.6, 1.65, 12.1, 5.0, bullets([
        f"Em 2026, os locais alagados votaram um pouco mais em Lula que o restante: diferença Lula − Flávio Bolsonaro de "
        f"{mp(dn['pct_lula'] - dn['pct_flavio'])} contra {mp(fo['pct_lula'] - fo['pct_flavio'])} pontos.",
        f"A distância entre os grupos já existia em 2022 e em 2018 (entre {mp(min(dist18))} e {mp(max(dist18))} pontos). De 2022 "
        f"para 2026, os dois grupos andaram quase o mesmo para o PL ({mp(md['mudanca_margem'])} e "
        f"{mp(mf['mudanca_margem'])} pontos), com qualquer definição de área alagada e sem padrão entre as cidades: "
        f"a diferença, {mp0(mi['voto_22_26']['estimativa'])} ponto, fica {margem(mi['voto_22_26'])}.",
        f"O que mudou mais nos locais alagados foi o eleitorado: {mp(md['mudanca_aptos_pct'])}% de eleitores aptos "
        f"(contra {mp(mf['mudanca_aptos_pct'])}%, {margem(mi['aptos_22_26'])}) e {mp(md['mudanca_abstencao'])} pontos "
        f"de abstenção (contra {mp(mf['mudanca_abstencao'])}, {margem(mi['abstencao_22_26'])}). No conjunto, de 2018 "
        "para 2022, antes da enchente, os dois grupos andavam juntos; em Canoas, os alagados já cresciam menos.",
        "Brancos e nulos subiram igual nos dois grupos.",
    ], tam=18, antes=16))
    s = d.novo("Fechamento", "Limites e fontes")
    texto(s, 0.6, 1.65, 6.0, 5.2, [par("Limites", 13, TXT2, negrito=True)] + bullets([
        "A unidade é o local de votação, não a casa do eleitor: o TSE não diz onde cada um mora.",
        "A mancha de satélite subestima a água entre prédios; o resultado foi refeito com outras manchas, com folga "
        "de 500 m e com peso pela população.",
        "Os candidatos e o cenário político mudam de uma eleição para outra (Haddad e Lula pelo PT; Jair Bolsonaro "
        "pelo PSL e pelo PL; Flávio Bolsonaro pelo PL): o que se compara é a distância entre os grupos, não o nível "
        "de cada ano.",
        f"{N['excluidos_sem_coord']['locais']} locais sem coordenada no TSE ficaram fora dos mapas de 2026 "
        f"({pct(N['excluidos_sem_coord']['pct_validos'])} dos votos válidos).",
        "O intervalo de 95% mede só o acaso de quais locais e zonas existem; não cobre erros da mancha nem outras "
        "diferenças entre as áreas.",
        "Comparação descritiva: não mede o efeito da enchente sobre o voto.",
    ], tam=13, antes=8))
    texto(s, 7.0, 1.65, 5.8, 5.2, [par("Fontes", 13, TXT2, negrito=True)] + bullets([
        "TSE: votação por seção e eleitorado por local de votação, 1º turno de 2018, 2022 e 2026.",
        "UFRGS/IPH, Possantti et al.: Banco de dados das cheias da RH do Guaíba v1.4, área diretamente afetada em 6 e "
        "8/5/2024 e simulação do Guaíba a 5,30 m (CC BY 4.0).",
        "Governo do RS/SPGG: área inundada em 23/5/2024 (checagem).",
        "IBGE: malha municipal 2025 e setores do Censo 2022.",
        "SEMA-RS: base cartográfica 1:25.000.",
    ], tam=13, antes=8))

    out = SAIDA / f"apresentacao_{ESTILO}.pptx"
    d.prs.save(out)
    return out, d.n


def main():
    L, N, NM = MM.carregar()
    vmax = float(L["m_locais"][["validos22", "validos26"]].max().max())
    N18 = json.loads((SAIDA / "numeros_2018.json").read_text(encoding="utf-8"))
    NI = json.loads((SAIDA / "numeros_incerteza.json").read_text(encoding="utf-8"))
    F = figuras(L, N, NM, N18, NI, vmax)
    out, n = montar(L, N, NM, N18, NI, F)
    print(out, n, "slides")


if __name__ == "__main__":
    main()
